import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}


class UnsupportedDocumentError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None
    text: str


@dataclass(frozen=True)
class TextChunk:
    chunk_index: int
    page_number: int | None
    content: str


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pages(filename: str, data: bytes) -> list[ExtractedPage]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError("Supported file types are PDF, DOCX, TXT, and CSV.")
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise UnsupportedDocumentError("The uploaded file is not a valid PDF.")
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            pages = [
                ExtractedPage(index, clean_text(page.extract_text() or ""))
                for index, page in enumerate(reader.pages, 1)
            ]
        except (PdfReadError, OSError, ValueError) as error:
            raise UnsupportedDocumentError("The PDF could not be read. Check that it is not encrypted or damaged.") from error
    elif suffix == ".docx":
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise UnsupportedDocumentError("The uploaded file is not a valid DOCX document.")
        try:
            document = DocxDocument(io.BytesIO(data))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as error:
            raise UnsupportedDocumentError("The DOCX document could not be read.") from error
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        pages = [ExtractedPage(None, clean_text("\n".join(lines)))]
    elif suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
            rows = list(csv.reader(io.StringIO(text)))
        except (UnicodeDecodeError, csv.Error) as error:
            raise UnsupportedDocumentError("The CSV file must be valid UTF-8 text.") from error
        pages = [ExtractedPage(None, clean_text("\n".join(" | ".join(row) for row in rows)))]
    else:
        try:
            pages = [ExtractedPage(None, clean_text(data.decode("utf-8-sig")))]
        except UnicodeDecodeError as error:
            raise UnsupportedDocumentError("Text documents must be valid UTF-8.") from error
    pages = [page for page in pages if page.text]
    if not pages:
        raise EmptyDocumentError("No extractable text was found in the uploaded document.")
    return pages


def chunk_pages(pages: list[ExtractedPage], chunk_size: int = 900, overlap: int = 140) -> list[TextChunk]:
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Chunk size must be positive and overlap must be smaller than the chunk size.")

    result: list[TextChunk] = []
    for page in pages:
        text = page.text.strip()
        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
                if boundary > start + chunk_size // 2:
                    end = boundary + (1 if text[boundary] == "\n" else 2)
            content = text[start:end].strip()
            if content:
                result.append(TextChunk(len(result), page.page_number, content))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return result
